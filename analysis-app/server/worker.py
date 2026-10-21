import sys,json,io,zipfile
from pathlib import Path

def safe_zip(data):
    z=zipfile.ZipFile(io.BytesIO(data))
    if sum(x.file_size for x in z.infolist())>64*1024*1024: raise ValueError('ไฟล์ขยายใหญ่เกินกำหนด')
    return z

def extract(p):
    data=Path(p['path']).read_bytes(); ext=p['extension'].lower(); pages=[]; truncated=False
    if ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(io.BytesIO(data))
        if len(reader.pages)>200: raise ValueError('รับ PDF ไม่เกิน 200 หน้า')
        for i,page in enumerate(reader.pages):
            raw=page.extract_text() or ''; truncated=truncated or len(raw)>25000
            pages.append({'page':i+1,'text':raw[:25000]})
    elif ext=='.docx' and data[:2]==b'PK':
        safe_zip(data)
        from docx import Document
        d=Document(io.BytesIO(data)); text='\n'.join(x.text for x in d.paragraphs)
        text+='\n'+'\n'.join(' | '.join(c.text for c in r.cells) for t in d.tables for r in t.rows)
        truncated=len(text)>100000;pages=[{'page':1,'text':text[:100000]}]
    elif ext in ('.txt','.md','.csv') or (ext=='.docx' and data[:2]!=b'PK'):
        try:text=data.decode('utf-8-sig')
        except UnicodeDecodeError:text=data.decode('cp874')
        truncated=len(text)>100000;pages=[{'page':1,'text':text[:100000]}]
    elif ext=='.xlsx':
        safe_zip(data)
        from openpyxl import load_workbook
        w=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
        for i,s in enumerate(w):
            if s.max_row and s.max_row>5000: truncated=True
            text='\n'.join(f'{s.title}!{r[0].row}: '+' | '.join(str(c.value if c.value is not None else '') for c in r) for r in s.iter_rows(max_row=5000) if any(c.value is not None for c in r))
            truncated=truncated or len(text)>100000
            pages.append({'page':i+1,'sheet':s.title,'text':text[:100000]})
    text='\n\n'.join(f'หน้า {x["page"]}\n{x["text"]}' for x in pages)
    truncated=truncated or len(text)>160000;text=text[:160000]
    return {'pages':pages,'text':text,'truncated':truncated,'reviewed':False,'status':'EXTRACTED_NEEDS_REVIEW' if text.strip() else 'UNRESOLVED_NEEDS_TRANSCRIPTION'}

def export(p):
    from docx import Document
    from docx.shared import Cm,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d=Document(); sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        s=d.styles[name];s.font.name='Tahoma';s.font.size=Pt(12 if name=='Normal' else 16 if name=='Title' else 13);s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_after=Pt(8);s.paragraph_format.line_spacing=1.25
        s.element.rPr.rFonts.set(qn('w:eastAsia'),'Tahoma');s.element.rPr.rFonts.set(qn('w:cs'),'Tahoma')
    for style in d.styles:
        if style.element.pPr is not None:
            for border in list(style.element.pPr.findall(qn('w:pBdr'))):style.element.pPr.remove(border)
    for paragraph in d.paragraphs:
        if paragraph._p.pPr is not None:
            for border in list(paragraph._p.pPr.findall(qn('w:pBdr'))):paragraph._p.pPr.remove(border)
    d.add_paragraph('ร่างคำวินิจฉัยเบื้องต้น',style='Title')
    d.add_paragraph(p['title'])
    d.add_paragraph('เอกสารสำหรับตรวจและแก้ไขโดยผู้รับผิดชอบ ข้อความนี้ยังไม่ใช่คำวินิจฉัยทางราชการ')
    d.add_paragraph(f'รหัสงานวิเคราะห์ {p["job_id"]}\nฉบับที่เลือก {p["revision_id"]}\nรุ่นฐานอ้างอิง {p["bundle_version"]}')
    if p.get('external_case_ref'): d.add_paragraph('เลขอ้างอิงจากภายนอก '+p['external_case_ref'])
    for title,text in p['sections'].items():
        d.add_heading(title,level=1)
        for line in text.split('\n'): d.add_paragraph(line)
    d.add_heading('รายการอ้างอิงของฉบับนี้',level=1)
    if not p['references']: d.add_paragraph('ยังไม่มีรายการอ้างอิงที่ยืนยัน')
    for r in p['references']:
        d.add_paragraph(f'[{r["id"]}] {r.get("label",r["id"])} | {r.get("source_pdf","เอกสารของงาน")} | หน้า {r.get("page","ตรวจข้อความ")} | '+('ตรวจและอนุมัติแล้ว' if r.get('reviewed') else 'รอตรวจ'))
    d.add_heading('ข้อมูลที่ยังรอตรวจ',level=1)
    for gap in p.get('gaps',[]):d.add_paragraph(gap)
    if not p.get('gaps'):d.add_paragraph('ไม่มีรายการช่องว่างที่บันทึกในฉบับนี้ ผู้ใช้ต้องตรวจเนื้อหาก่อนใช้งาน')
    footer=sec.footer.paragraphs[0]; footer.alignment=2
    footer.add_run('ร่างเบื้องต้น • หน้า ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    d.save(p['path']);return {'path':p['path']}

if __name__=='__main__':
    try:
        payload=json.loads(sys.stdin.buffer.read().decode('utf-8')); out=extract(payload) if sys.argv[1]=='extract' else export(payload)
        print(json.dumps(out,ensure_ascii=True))
    except Exception as e:
        print(json.dumps({'error':str(e)},ensure_ascii=True));sys.exit(1)
